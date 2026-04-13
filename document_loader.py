import os
import magic
from typing import List
from langchain_core.documents import Document


class DocumentLoader:
    def __init__(self):
        self.mime = magic.Magic(mime=True)

    def load_document(self, file_path: str) -> List[Document]:
        """加载文档并返回文档对象列表"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")

        file_extension = os.path.splitext(file_path)[1].lower()

        # 明确检查文件扩展名
        if file_extension == ".pdf":
            return self._load_pdf(file_path)
        elif file_extension == ".txt":
            return self._load_txt(file_path)
        elif file_extension == ".docx":
            return self._load_docx(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {file_extension}")

    def _load_pdf(self, file_path: str) -> List[Document]:
        """加载 PDF 文档"""
        try:
            import pypdf

            documents = []
            with open(file_path, "rb") as file:
                reader = pypdf.PdfReader(file)
                for page_num, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        doc = Document(
                            page_content=text,
                            metadata={
                                "source": file_path,
                                "page": page_num + 1,
                                "total_pages": len(reader.pages),
                            },
                        )
                        documents.append(doc)
            return documents
        except ImportError:
            raise ImportError("需要安装 pypdf 来处理 PDF 文件")

    def _load_txt(self, file_path: str) -> List[Document]:
        """加载 TXT 文档"""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            text = file.read()
            doc = Document(page_content=text, metadata={"source": file_path})
            return [doc]

    def _load_docx(self, file_path: str) -> List[Document]:
        """加载 DOCX 文档"""
        try:
            from docx import Document as DocxDocument

            docx = DocxDocument(file_path)
            text = "\n".join([para.text for para in docx.paragraphs])
            doc = Document(page_content=text, metadata={"source": file_path})
            return [doc]
        except ImportError:
            raise ImportError("需要安装 python-docx 来处理 DOCX 文件")
